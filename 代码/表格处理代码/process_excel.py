import os
import argparse
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

def process_excel(input_path, output_path):
    wb = load_workbook(input_path)
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        headers = []
        col_indices = []
        
        for col in range(1, ws.max_column + 1):
            header = ws.cell(row=1, column=col).value
            headers.append(header)
            if header in ['post_id', 'comment_id', 'comment_like_count']:
                col_indices.append(col)
        
        col_indices.sort(reverse=True)
        
        for col_idx in col_indices:
            ws.delete_cols(col_idx)
        
        headers = [h for h in headers if h not in ['post_id', 'comment_id', 'comment_like_count']]
        
        pics_col_idx = None
        for idx, header in enumerate(headers):
            if header == 'pics':
                pics_col_idx = idx + 1
                break
        
        if pics_col_idx:
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=pics_col_idx)
                cell_value = cell.value
                
                if cell_value is None or cell_value == '' or str(cell_value).strip() == '':
                    cell.value = '❌️'
                else:
                    cell.value = '√'
    
    wb.save(output_path)
    print(f"处理完成！结果已保存到: {output_path}")

def batch_process_excel(folder_path):
    if not os.path.exists(folder_path):
        print(f"错误：目录不存在: {folder_path}")
        return
    
    excel_files = [f for f in os.listdir(folder_path) if f.endswith('.xlsx') and not f.endswith('_processed.xlsx')]
    
    if not excel_files:
        print("未找到需要处理的Excel文件")
        return
    
    print(f"找到 {len(excel_files)} 个Excel文件待处理...")
    
    for idx, filename in enumerate(excel_files, 1):
        input_path = os.path.join(folder_path, filename)
        name_without_ext = os.path.splitext(filename)[0]
        output_path = os.path.join(folder_path, f"{name_without_ext}_processed.xlsx")
        
        print(f"\n[{idx}/{len(excel_files)}] 正在处理: {filename}")
        process_excel(input_path, output_path)
    
    print(f"\n批量处理完成！共处理 {len(excel_files)} 个文件")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description='处理Excel文件：删除指定列并替换pics列内容')
    parser.add_argument('path', nargs='?', default=None, 
                        help='要处理的Excel文件路径或包含Excel文件的目录路径')
    
    args = parser.parse_args()
    
    if args.path:
        if os.path.isfile(args.path) and args.path.endswith('.xlsx'):
            name_without_ext = os.path.splitext(args.path)[0]
            output_path = f"{name_without_ext}_processed.xlsx"
            print(f"正在处理文件: {args.path}")
            process_excel(args.path, output_path)
        elif os.path.isdir(args.path):
            batch_process_excel(args.path)
        else:
            print(f"错误：无效的路径: {args.path}")
            print("请提供有效的Excel文件路径或目录路径")
    else:
        default_folder = r'c:\Users\ThinkPad\Desktop\weibopachong\scan_details'
        print(f"未指定路径，使用默认目录: {default_folder}")
        batch_process_excel(default_folder)